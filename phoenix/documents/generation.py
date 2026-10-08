"""Caller-owned in-memory reports. Artifacts are untrusted data, never saved files.

Native rendering/compression is synchronous and cannot be hard-preempted. PDF V1
uses Base-14 fonts and accepts only cp1252 text. No tool or permission authority.
"""

import asyncio
import hashlib
import logging
import re
import threading
from collections.abc import Iterator
from contextlib import contextmanager
from enum import StrEnum
from io import BytesIO
from typing import Annotated, Literal, NoReturn, Self
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import docx
import pypdf
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph as WordParagraph
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    ModelWrapValidatorHandler,
    field_validator,
    model_validator,
)
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from typing_extensions import Buffer


class DocumentFormat(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    MARKDOWN = "markdown"


class DocumentGenerationErrorCode(StrEnum):
    INVALID_INPUT = "invalid_input"
    UNSUPPORTED_CHARACTER = "unsupported_character"
    LIMIT_EXCEEDED = "limit_exceeded"
    GENERATION_FAILED = "generation_failed"
    VERIFICATION_FAILED = "verification_failed"
    CANCELLED = "cancelled"
    BUSY = "busy"
    CLOSED = "closed"
    WRONG_EVENT_LOOP = "wrong_event_loop"


class DocumentGenerationError(Exception):
    """Only a closed code, without retained content or native diagnostics."""

    def __init__(self, code: DocumentGenerationErrorCode) -> None:
        if not any(code is member for member in DocumentGenerationErrorCode):
            raise ValueError("invalid_input")
        self.code = code
        super().__init__(code.value)


def _fail(code: DocumentGenerationErrorCode) -> NoReturn:
    raise DocumentGenerationError(code)


def _plain(value: object, *, meaningful: bool = False) -> str:
    if type(value) is not str:
        _fail(DocumentGenerationErrorCode.INVALID_INPUT)
    try:
        value.encode("utf-8", errors="strict")
    except UnicodeError:
        _fail(DocumentGenerationErrorCode.INVALID_INPUT)
    if any((ord(c) < 32 and c not in "\t\n\r") or 127 <= ord(c) <= 159 for c in value):
        _fail(DocumentGenerationErrorCode.INVALID_INPUT)
    if meaningful and not value.strip():
        _fail(DocumentGenerationErrorCode.INVALID_INPUT)
    return value


_Text = Annotated[str, Field(strict=True)]
_Positive = Annotated[int, Field(strict=True, gt=0)]
_MEDIA = {
    DocumentFormat.PDF: "application/pdf",
    DocumentFormat.DOCX: (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
    DocumentFormat.MARKDOWN: "text/markdown; charset=utf-8",
}
_EXT = {
    DocumentFormat.PDF: ".pdf",
    DocumentFormat.DOCX: ".docx",
    DocumentFormat.MARKDOWN: ".md",
}


class _Contract(BaseModel):
    model_config = ConfigDict(
        strict=True,
        frozen=True,
        extra="forbid",
        revalidate_instances="always",
        validate_default=True,
        hide_input_in_errors=True,
    )

    @model_validator(mode="wrap")
    @classmethod
    def exact(cls, value: object, handler: ModelWrapValidatorHandler[Self]) -> Self:
        try:
            if isinstance(value, BaseModel):
                if type(value) is not cls:
                    _fail(DocumentGenerationErrorCode.INVALID_INPUT)
                value = {**vars(value), **(value.__pydantic_extra__ or {})}
            if type(value) is not dict:
                _fail(DocumentGenerationErrorCode.INVALID_INPUT)
            return handler(value)
        except DocumentGenerationError:
            raise
        except Exception:
            raise DocumentGenerationError(
                DocumentGenerationErrorCode.INVALID_INPUT
            ) from None

    @model_validator(mode="after")
    def plain_fields(self) -> Self:
        for value in vars(self).values():
            if type(value) is str:
                _plain(value)
        return self


class ParagraphBlock(_Contract):
    kind: Literal["paragraph"]
    text: _Text


class BulletListBlock(_Contract):
    kind: Literal["bullet_list"]
    items: Annotated[tuple[_Text, ...], Field(min_length=1)]

    @field_validator("items")
    @classmethod
    def item_text(cls, value: tuple[str, ...]) -> tuple[str, ...]:
        return tuple(_plain(item) for item in value)


class TableBlock(_Contract):
    kind: Literal["table"]
    headers: Annotated[tuple[_Text, ...], Field(min_length=1)]
    rows: tuple[tuple[_Text, ...], ...]

    @model_validator(mode="after")
    def width(self) -> Self:
        for text in self.headers:
            _plain(text)
        for row in self.rows:
            if len(row) != len(self.headers):
                _fail(DocumentGenerationErrorCode.INVALID_INPUT)
            for text in row:
                _plain(text)
        return self


ReportBlock = Annotated[
    ParagraphBlock | BulletListBlock | TableBlock, Field(discriminator="kind")
]


class ReportSection(_Contract):
    heading: _Text
    blocks: Annotated[tuple[ReportBlock, ...], Field(min_length=1)]

    @field_validator("heading", mode="before")
    @classmethod
    def meaningful(cls, value: object) -> str:
        return _plain(value, meaningful=True)


class ReportProvenance(_Contract):
    label: _Text
    reference: _Text

    @field_validator("label", "reference", mode="before")
    @classmethod
    def meaningful(cls, value: object) -> str:
        return _plain(value, meaningful=True)


class ReportDocument(_Contract):
    document_id: _Text
    title: _Text
    subtitle: _Text | None = None
    author: _Text | None = None
    sections: Annotated[tuple[ReportSection, ...], Field(min_length=1)]
    provenance: tuple[ReportProvenance, ...] = ()

    @field_validator("document_id", "title", mode="before")
    @classmethod
    def meaningful(cls, value: object) -> str:
        return _plain(value, meaningful=True)

    @field_validator("document_id")
    @classmethod
    def identity_bound(cls, value: str) -> str:
        if len(value) > 128:
            _fail(DocumentGenerationErrorCode.INVALID_INPUT)
        return value


def _filename(value: object, fmt: DocumentFormat) -> str:
    text = _plain(value, meaningful=True)
    if (
        not 1 <= len(text) <= 255
        or text != text.strip()
        or any(c in text for c in "/\\:")
        or any(ord(c) < 32 or 127 <= ord(c) <= 159 for c in text)
        or text in (".", "..")
        or not text.lower().endswith(_EXT[fmt])
        or not text[: -len(_EXT[fmt])].strip(".")
        or ".." in text
    ):
        _fail(DocumentGenerationErrorCode.INVALID_INPUT)
    return text


class DocumentGenerationRequest(_Contract):
    generation_id: _Text
    format: DocumentFormat
    suggested_filename: _Text
    document: ReportDocument

    @field_validator("generation_id", mode="before")
    @classmethod
    def meaningful(cls, value: object) -> str:
        return _plain(value, meaningful=True)

    @model_validator(mode="after")
    def filename(self) -> Self:
        _filename(self.suggested_filename, self.format)
        return self


class DocumentGenerationLimits(_Contract):
    max_total_input_chars: _Positive
    max_sections: _Positive
    max_blocks: _Positive
    max_list_items: _Positive
    max_table_rows: _Positive
    max_table_columns: _Positive
    max_cell_chars: _Positive
    max_provenance_items: _Positive
    max_output_bytes: _Positive
    max_pdf_pages: _Positive


class GeneratedArtifact(_Contract):
    generation_id: _Text
    document_id: _Text
    format: DocumentFormat
    suggested_filename: _Text
    media_type: _Text
    artifact_sha256: _Text
    byte_size: _Positive
    page_count: _Positive | None
    verified: Annotated[bool, Field(strict=True)]
    artifact_bytes: Annotated[bytes, Field(strict=True, min_length=1)] = Field(
        repr=False, exclude=True
    )

    @field_validator("artifact_bytes", mode="before")
    @classmethod
    def bytes_exact(cls, value: object) -> bytes:
        if type(value) is not bytes or not value:
            _fail(DocumentGenerationErrorCode.INVALID_INPUT)
        return value

    @model_validator(mode="after")
    def evidence(self) -> Self:
        _plain(self.generation_id, meaningful=True)
        _plain(self.document_id, meaningful=True)
        if len(self.document_id) > 128:
            _fail(DocumentGenerationErrorCode.INVALID_INPUT)
        _filename(self.suggested_filename, self.format)
        if (
            self.media_type != _MEDIA[self.format]
            or self.artifact_sha256 != hashlib.sha256(self.artifact_bytes).hexdigest()
            or self.byte_size != len(self.artifact_bytes)
            or not self.verified
            or (self.format is DocumentFormat.PDF) != (self.page_count is not None)
        ):
            _fail(DocumentGenerationErrorCode.INVALID_INPUT)
        return self


def _document_text(document: ReportDocument) -> Iterator[str]:
    yield document.document_id
    yield document.title
    for optional in (document.subtitle, document.author):
        if optional is not None:
            yield optional
    for section in document.sections:
        yield section.heading
        for block in section.blocks:
            if isinstance(block, ParagraphBlock):
                yield block.text
            elif isinstance(block, BulletListBlock):
                yield from block.items
            else:
                yield from block.headers
                for row in block.rows:
                    yield from row
    for item in document.provenance:
        yield item.label
        yield item.reference


def _admit(
    request: DocumentGenerationRequest, limits: DocumentGenerationLimits
) -> None:
    document = request.document
    total = len(request.generation_id) + len(request.suggested_filename)
    total += sum(len(text) for text in _document_text(document))
    blocks = lists = rows = 0
    exceeded = len(document.sections) > limits.max_sections
    for section in document.sections:
        blocks += len(section.blocks)
        for block in section.blocks:
            if isinstance(block, BulletListBlock):
                lists += len(block.items)
            elif isinstance(block, TableBlock):
                rows += len(block.rows)
                exceeded |= len(block.headers) > limits.max_table_columns
                exceeded |= any(
                    len(cell) > limits.max_cell_chars for cell in block.headers
                )
                exceeded |= any(
                    len(cell) > limits.max_cell_chars
                    for row in block.rows
                    for cell in row
                )
    if exceeded or (
        total > limits.max_total_input_chars
        or blocks > limits.max_blocks
        or lists > limits.max_list_items
        or rows > limits.max_table_rows
        or len(document.provenance) > limits.max_provenance_items
    ):
        _fail(DocumentGenerationErrorCode.LIMIT_EXCEEDED)
    if request.format is DocumentFormat.PDF:
        try:
            # Only rendered text, not opaque IDs or the display filename.
            for text in list(_document_text(document))[1:]:
                text.encode("cp1252", errors="strict")
        except UnicodeError:
            _fail(DocumentGenerationErrorCode.UNSUPPORTED_CHARACTER)


def _lines(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _pdf_text(text: str) -> str:
    return escape(_lines(text)).replace("\n", "<br/>")


def _markdown_text(text: str) -> str:
    # Flatten user line breaks into inert literal backslash-n. Structure is AST-only.
    escaped = re.sub(r"([\\`*_{}\[\]()#+.!|<>~\-=])", r"\\\1", _lines(text))
    return escaped.replace("\n", r"\n")


class _Output(BytesIO):
    def __init__(self, maximum: int) -> None:
        super().__init__()
        self.maximum = maximum

    def write(self, data: Buffer, /) -> int:
        data = bytes(data)
        if self.tell() + len(data) > self.maximum:
            _fail(DocumentGenerationErrorCode.LIMIT_EXCEEDED)
        return super().write(data)


class _PdfPageLimitExceeded(Exception):
    """Private callback signal; ReportLab may reconstruct it using a string."""


# A minimal, code-owned empty OOXML package prevents docx.Document(None) and lazy
# default styles from opening bundled template files. No user text enters this XML.
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/package/2006/relationships"
_OR = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
_DOCX_PARTS = {
    "[Content_Types].xml": (
        f'<Types xmlns="{_CT}"><Default Extension="rels" ContentType="'
        f'application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="applica'
        "tion/vnd.openxmlformats-officedocument.wordprocessingml.docu"
        'ment.main+xml"/>'
        '<Override PartName="/word/styles.xml" ContentType="applicati'
        "on/vnd.openxmlformats-officedocument.wordprocessingml.styles"
        '+xml"/>'
        '<Override PartName="/word/numbering.xml" ContentType="applic'
        "ation/vnd.openxmlformats-officedocument.wordprocessingml.num"
        'bering+xml"/></Types>'
    ),
    "_rels/.rels": (
        f'<Relationships xmlns="{_R}"><Relationship Id="rId1" Type="{_OR}'
        f'officeDocument" Target="word/document.xml"/></Relationships>'
    ),
    "word/document.xml": (
        f'<w:document xmlns:w="{_W}"><w:body><w:sectPr><w:pgSz w:w="11'
        f'906" w:h="16838"/><w:pgMar w:top="1440" w:right="1440" w:bot'
        f'tom="1440" w:left="1440"/></w:sectPr></w:body></w:document>'
    ),
    "word/_rels/document.xml.rels": (
        f'<Relationships xmlns="{_R}"><Relationship Id="rId1" Type="{_OR}'
        f'styles" Target="styles.xml"/><Relationship Id="rId2" Type="{_OR}'
        f'numbering" Target="numbering.xml"/></Relationships>'
    ),
    "word/styles.xml": (
        f'<w:styles xmlns:w="{_W}">'
        '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"'
        '><w:name w:val="Normal"/><w:rPr><w:rFonts w:ascii="Calibri" '
        'w:hAnsi="Calibri"/><w:sz w:val="22"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val='
        '"Title"/><w:basedOn w:val="Normal"/><w:rPr><w:sz w:val="44"/'
        "><w:b/></w:rPr></w:style>"
        '<w:style w:type="paragraph" w:styleId="Subtitle"><w:name w:v'
        'al="Subtitle"/><w:basedOn w:val="Normal"/><w:rPr><w:i/></w:r'
        "Pr></w:style>"
        '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:v'
        'al="heading 1"/><w:basedOn w:val="Normal"/><w:pPr><w:keepNex'
        't/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:sz w:val'
        '="30"/></w:rPr></w:style>'
        '<w:style w:type="paragraph" w:styleId="ListBullet"><w:name w'
        ':val="List Bullet"/><w:basedOn w:val="Normal"/><w:pPr><w:num'
        'Pr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr>'
        "</w:style>"
        '<w:style w:type="table" w:default="1" w:styleId="TableNormal'
        '"><w:name w:val="Normal Table"/></w:style>'
        '<w:style w:type="table" w:styleId="TableGrid"><w:name w:val='
        '"Table Grid"/><w:basedOn w:val="TableNormal"/><w:tblPr><w:tb'
        "lBorders>"
        + "".join(
            f'<w:{edge} w:val="single" w:sz="4" w:color="808080"/>'
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
        )
        + "</w:tblBorders></w:tblPr></w:style></w:styles>"
    ),
    "word/numbering.xml": (
        f'<w:numbering xmlns:w="{_W}"><w:abstractNum w:abstractNumId="'
        f'0"><w:multiLevelType w:val="singleLevel"/><w:lvl w:ilvl="0">'
        f'<w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:v'
        f'al="•"/><w:lvlJc w:val="left"/><w:pPr><w:tabs><w:tab w:val="'
        f'num" w:pos="720"/></w:tabs><w:ind w:left="720" w:hanging="36'
        f'0"/></w:pPr></w:lvl></w:abstractNum><w:num w:numId="1"><w:ab'
        f'stractNumId w:val="0"/></w:num></w:numbering>'
    ),
}


def _zip_write(archive: ZipFile, name: str, data: bytes) -> None:
    info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
    info.compress_type = ZIP_DEFLATED
    info.create_system = 0
    archive.writestr(info, data)


@contextmanager
def _private_pdf_logs() -> Iterator[None]:
    factory = logging.getLogRecordFactory()
    owner = threading.get_ident()

    def private(*args: object, **kwargs: object) -> logging.LogRecord:
        record = factory(*args, **kwargs)
        if record.thread == owner and record.name.startswith("pypdf"):
            record.msg, record.args = "document verification diagnostic", ()
            record.exc_info = record.exc_text = record.stack_info = None
        return record

    logging.setLogRecordFactory(private)
    try:
        yield
    finally:
        logging.setLogRecordFactory(factory)


class DocumentGenerator:
    """One loop and one caller-owned generation; no background resources."""

    def __init__(self, limits: DocumentGenerationLimits) -> None:
        self._limits = DocumentGenerationLimits.model_validate(limits)
        self._loop: asyncio.AbstractEventLoop | None = None
        self._active: str | None = None
        self._closed = False

    @property
    def active_generation_id(self) -> str | None:
        return self._active

    def _bind(self) -> None:
        loop = asyncio.get_running_loop()
        if self._loop is not None and self._loop is not loop:
            _fail(DocumentGenerationErrorCode.WRONG_EVENT_LOOP)
        self._loop = loop

    def _check(self) -> None:
        task = asyncio.current_task()
        if task is not None and task.cancelling():
            raise asyncio.CancelledError
        if self._closed:
            _fail(DocumentGenerationErrorCode.CLOSED)

    async def _checkpoint(self) -> None:
        self._check()
        await asyncio.sleep(0)
        self._check()

    async def aclose(self) -> None:
        self._bind()
        self._closed = True

    async def generate(self, request: DocumentGenerationRequest) -> GeneratedArtifact:
        self._bind()
        self._check()
        if self._active is not None:
            _fail(DocumentGenerationErrorCode.BUSY)
        request = DocumentGenerationRequest.model_validate(request)
        _admit(request, self._limits)
        self._active = request.generation_id
        try:
            await self._checkpoint()
            try:
                if request.format is DocumentFormat.PDF:
                    data = await self._pdf(request.document)
                elif request.format is DocumentFormat.DOCX:
                    data = await self._docx(request.document)
                else:
                    data = await self._markdown(request.document)
            except DocumentGenerationError:
                self._check()
                raise
            except Exception:
                self._check()
                raise DocumentGenerationError(
                    DocumentGenerationErrorCode.GENERATION_FAILED
                ) from None
            await self._checkpoint()
            pages = await self._verify(request, data)
            await self._checkpoint()
            result = GeneratedArtifact(
                generation_id=request.generation_id,
                document_id=request.document.document_id,
                format=request.format,
                suggested_filename=request.suggested_filename,
                media_type=_MEDIA[request.format],
                artifact_sha256=hashlib.sha256(data).hexdigest(),
                byte_size=len(data),
                page_count=pages,
                verified=True,
                artifact_bytes=data,
            )
            result = GeneratedArtifact.model_validate(result)
            self._check()
            return result
        finally:
            self._active = None

    async def _markdown(self, document: ReportDocument) -> bytes:
        with _Output(self._limits.max_output_bytes) as output:

            def line(text: str = "") -> None:
                output.write((text + "\n").encode("utf-8"))

            line("# " + _markdown_text(document.title))
            line()
            for optional in (document.subtitle, document.author):
                if optional is not None:
                    line(_markdown_text(optional))
                    line()
            for section in document.sections:
                await self._checkpoint()
                line("## " + _markdown_text(section.heading))
                line()
                for block in section.blocks:
                    await self._checkpoint()
                    if isinstance(block, ParagraphBlock):
                        line(_markdown_text(block.text))
                    elif isinstance(block, BulletListBlock):
                        for item in block.items:
                            await self._checkpoint()
                            line("- " + _markdown_text(item))
                    else:
                        line(
                            "| "
                            + " | ".join(_markdown_text(x) for x in block.headers)
                            + " |"
                        )
                        line("| " + " | ".join("---" for _ in block.headers) + " |")
                        for row in block.rows:
                            await self._checkpoint()
                            line(
                                "| " + " | ".join(_markdown_text(x) for x in row) + " |"
                            )
                    line()
            if document.provenance:
                line("## Sources")
                line()
                for source in document.provenance:
                    await self._checkpoint()
                    line(
                        "- "
                        + _markdown_text(source.label)
                        + " — "
                        + _markdown_text(source.reference)
                    )
            return output.getvalue()

    async def _pdf(self, document: ReportDocument) -> bytes:
        styles = {
            name: ParagraphStyle(
                name, fontName=font, fontSize=size, leading=size * 1.35, spaceAfter=8
            )
            for name, font, size in (
                ("title", "Helvetica-Bold", 22),
                ("heading", "Helvetica-Bold", 14),
                ("body", "Helvetica", 10),
            )
        }
        story: list[object] = []

        def paragraph(text: str, style: str = "body") -> Paragraph:
            return Paragraph(_pdf_text(text), styles[style])

        story.append(paragraph(document.title, "title"))
        for optional in (document.subtitle, document.author):
            if optional is not None:
                story.append(paragraph(optional))
        for section in document.sections:
            await self._checkpoint()
            story.append(paragraph(section.heading, "heading"))
            for block in section.blocks:
                await self._checkpoint()
                if isinstance(block, ParagraphBlock):
                    story.append(paragraph(block.text))
                elif isinstance(block, BulletListBlock):
                    for item in block.items:
                        await self._checkpoint()
                        story.append(paragraph("• " + item))
                else:
                    cells = [[paragraph(cell) for cell in block.headers]]
                    for row in block.rows:
                        await self._checkpoint()
                        cells.append([paragraph(cell) for cell in row])
                    table = Table(
                        cells,
                        colWidths=[451 / len(block.headers)] * len(block.headers),
                        repeatRows=1,
                        hAlign="LEFT",
                    )
                    table.setStyle(
                        TableStyle(
                            [
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                (
                                    "BACKGROUND",
                                    (0, 0),
                                    (-1, 0),
                                    colors.HexColor("#eeeeee"),
                                ),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                            ]
                        )
                    )
                    story.extend((table, Spacer(1, 8)))
        if document.provenance:
            story.append(paragraph("Sources", "heading"))
            for source in document.provenance:
                await self._checkpoint()
                story.append(paragraph(source.label + " — " + source.reference))

        def footer(canvas: Canvas, _: SimpleDocTemplate) -> None:
            task = asyncio.current_task()
            if task is not None and task.cancelling():
                raise asyncio.CancelledError
            if canvas.getPageNumber() > self._limits.max_pdf_pages:
                raise _PdfPageLimitExceeded("page limit exceeded")
            canvas.saveState()
            canvas.setFont("Helvetica", 9)
            canvas.drawRightString(523, 36, "Page " + str(canvas.getPageNumber()))
            canvas.restoreState()

        await self._checkpoint()
        with _Output(self._limits.max_output_bytes) as output:
            template = SimpleDocTemplate(
                output,
                pagesize=(595, 842),
                leftMargin=72,
                rightMargin=72,
                topMargin=72,
                bottomMargin=60,
                invariant=1,
                pageCompression=1,
                title=document.title,
                author=document.author or "",
            )
            self._check()
            try:
                template.build(story, onFirstPage=footer, onLaterPages=footer)
            except _PdfPageLimitExceeded:
                self._check()
                raise DocumentGenerationError(
                    DocumentGenerationErrorCode.LIMIT_EXCEEDED
                ) from None
            self._check()
            return output.getvalue()

    async def _docx(self, document: ReportDocument) -> bytes:
        with BytesIO() as seed:
            with ZipFile(seed, "w") as archive:
                for name, text in _DOCX_PARTS.items():
                    _zip_write(archive, name, text.encode("utf-8"))
            seed.seek(0)
            self._check()
            word = docx.Document(seed)
            self._check()
        word.add_paragraph(document.title, "Title")
        for optional in (document.subtitle, document.author):
            if optional is not None:
                word.add_paragraph(optional, "Subtitle")
        for section in document.sections:
            await self._checkpoint()
            word.add_heading(section.heading, level=1)
            for block in section.blocks:
                await self._checkpoint()
                if isinstance(block, ParagraphBlock):
                    word.add_paragraph(block.text)
                elif isinstance(block, BulletListBlock):
                    for item in block.items:
                        await self._checkpoint()
                        word.add_paragraph(item, "List Bullet")
                else:
                    table = word.add_table(rows=1, cols=len(block.headers))
                    table.style = "Table Grid"
                    for cell, text in zip(
                        table.rows[0].cells, block.headers, strict=True
                    ):
                        cell.text = text
                    for row in block.rows:
                        await self._checkpoint()
                        for cell, text in zip(table.add_row().cells, row, strict=True):
                            cell.text = text
        if document.provenance:
            word.add_heading("Sources", level=1)
            for source in document.provenance:
                await self._checkpoint()
                word.add_paragraph(source.label + " — " + source.reference)
        await self._checkpoint()
        # Artifact acceptance bound: do not interrupt native ZIP serialization.
        # Ordinary BytesIO remains alive throughout save and archive cleanup.
        with BytesIO() as raw:
            word.save(raw)
            self._check()
            if len(raw.getvalue()) > self._limits.max_output_bytes:
                _fail(DocumentGenerationErrorCode.LIMIT_EXCEEDED)
            # python-docx ZIP timestamps are wall-clock based; canonicalize package
            # timestamps and order in memory for deterministic artifact bytes.
            raw.seek(0)
            with (
                ZipFile(raw) as archive,
                BytesIO() as output,
            ):
                with ZipFile(output, "w") as canonical:
                    for name in sorted(archive.namelist()):
                        await self._checkpoint()
                        _zip_write(canonical, name, archive.read(name))
                self._check()
                data = output.getvalue()
                if len(data) > self._limits.max_output_bytes:
                    _fail(DocumentGenerationErrorCode.LIMIT_EXCEEDED)
                return data

    async def _verify(
        self, request: DocumentGenerationRequest, data: bytes
    ) -> int | None:
        if len(data) > self._limits.max_output_bytes:
            _fail(DocumentGenerationErrorCode.LIMIT_EXCEEDED)
        await self._checkpoint()
        try:
            if request.format is DocumentFormat.PDF:
                if not data.startswith(b"%PDF-"):
                    _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
                with _private_pdf_logs(), BytesIO(data) as buffer:
                    reader = pypdf.PdfReader(buffer, strict=True)
                    self._check()
                    if reader.is_encrypted:
                        _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
                    count = len(reader.pages)
                    self._check()
                    if not 1 <= count <= self._limits.max_pdf_pages:
                        _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
                    extracted: list[str] = []
                    for page in reader.pages:
                        await self._checkpoint()
                        extracted.append(page.extract_text())
                        self._check()
                    # Layout wrapping changes whitespace, not title characters.
                    if " ".join(request.document.title.split()) not in " ".join(
                        "\n".join(extracted).split()
                    ):
                        _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
                return count
            if request.format is DocumentFormat.DOCX:
                if not data.startswith(b"PK\x03\x04"):
                    _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
                with BytesIO(data) as buffer, ZipFile(buffer) as archive:
                    names = archive.namelist()
                    if (
                        len(names) != len(set(names))
                        or not {"[Content_Types].xml", "word/document.xml"}
                        <= set(names)
                        or any(info.flag_bits & 1 for info in archive.infolist())
                    ):
                        _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
                await self._checkpoint()
                with BytesIO(data) as buffer:
                    word = docx.Document(buffer)
                    self._check()
                    if (
                        not word.paragraphs
                        or word.paragraphs[0].text != request.document.title
                    ):
                        _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
                    for block in word.iter_inner_content():
                        await self._checkpoint()
                        if isinstance(block, WordParagraph):
                            _ = block.text
                        elif isinstance(block, WordTable):
                            for row in block.rows:
                                await self._checkpoint()
                                for cell in row.cells:
                                    _ = cell.text
                return None
            text = data.decode("utf-8", errors="strict")
            if (
                "\0" in text
                or "\r" in text
                or not text.startswith(
                    "# " + _markdown_text(request.document.title) + "\n"
                )
                or not text.endswith("\n")
                or text.encode("utf-8") != data
            ):
                _fail(DocumentGenerationErrorCode.VERIFICATION_FAILED)
            return None
        except DocumentGenerationError:
            self._check()
            raise
        except Exception:
            self._check()
            raise DocumentGenerationError(
                DocumentGenerationErrorCode.VERIFICATION_FAILED
            ) from None
