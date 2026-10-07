"""Bounded in-memory adapters; content is untrusted and grants no authority.

PHX-032 owns admission, persistence, chunking and active ingestion state. PDF
native decompression/extraction is synchronous, not a hard memory/preemption
boundary. DOCX covers top-level paragraph/table text; images have no default OCR.
"""

import asyncio
import csv
import json
import logging
import re
import sys
import threading
import zipfile
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime
from enum import StrEnum
from io import BytesIO, StringIO
from typing import Annotated, NoReturn, Self, TypeVar
from xml.parsers import expat

import docx
import pypdf
from docx.table import Table
from docx.text.paragraph import Paragraph
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    ModelWrapValidatorHandler,
    field_validator,
    model_validator,
)

from phoenix.rag.ingestion import (
    DocumentParser,
    DocumentSource,
    ParsedDocument,
    ParsedSegment,
)


class UploadFormat(StrEnum):
    TEXT = "text"
    MARKDOWN = "markdown"
    CSV = "csv"
    CODE = "code"
    PDF = "pdf"
    DOCX = "docx"
    IMAGE = "image"


class UploadParserErrorCode(StrEnum):
    INVALID_INPUT = "invalid_input"
    UNSUPPORTED_FORMAT = "unsupported_format"
    MEDIA_TYPE_MISMATCH = "media_type_mismatch"
    SOURCE_TOO_LARGE = "source_too_large"
    DECODE_FAILURE = "decode_failure"
    MALFORMED_DOCUMENT = "malformed_document"
    ENCRYPTED_DOCUMENT = "encrypted_document"
    NO_EXTRACTABLE_TEXT = "no_extractable_text"
    LIMIT_EXCEEDED = "limit_exceeded"
    PARSER_UNAVAILABLE = "parser_unavailable"
    CANCELLED = "cancelled"
    WRONG_EVENT_LOOP = "wrong_event_loop"


class UploadParserError(Exception):
    """Closed category only; never retain library diagnostics or source text."""

    def __init__(self, code: UploadParserErrorCode) -> None:
        if not any(code is member for member in UploadParserErrorCode):
            raise ValueError("invalid_input")
        self.code = code
        super().__init__(code.value)


_CODE_EXTENSIONS = frozenset(
    ".py .pyi .js .jsx .ts .tsx .json .toml .yaml .yml .xml .html .htm .css "
    ".sql .sh .ps1 .bat .cmd .c .h .cpp .hpp .java .kt .go .rs .rb .php .cs".split()
)
_EXTENSIONS = {
    ".txt": UploadFormat.TEXT,
    ".md": UploadFormat.MARKDOWN,
    ".markdown": UploadFormat.MARKDOWN,
    ".csv": UploadFormat.CSV,
    ".pdf": UploadFormat.PDF,
    ".docx": UploadFormat.DOCX,
    ".png": UploadFormat.IMAGE,
    ".jpg": UploadFormat.IMAGE,
    ".jpeg": UploadFormat.IMAGE,
    ".webp": UploadFormat.IMAGE,
    **dict.fromkeys(_CODE_EXTENSIONS, UploadFormat.CODE),
}
_MEDIA = {
    UploadFormat.TEXT: "text/plain",
    UploadFormat.MARKDOWN: "text/markdown",
    UploadFormat.CSV: "text/csv",
    UploadFormat.CODE: "text/plain",
    UploadFormat.PDF: "application/pdf",
    UploadFormat.DOCX: (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ),
}
_IMAGE_MEDIA = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
}
_Positive = Annotated[int, Field(strict=True, gt=0)]
_T = TypeVar("_T")


def _fail(code: UploadParserErrorCode) -> NoReturn:
    raise UploadParserError(code)


def _text(value: object) -> str:
    if type(value) is not str or not value.strip():
        _fail(UploadParserErrorCode.INVALID_INPUT)
    text = value
    try:
        text.encode("utf-8")
    except UnicodeError:
        _fail(UploadParserErrorCode.INVALID_INPUT)
    return text


def _extension(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower()


class _Contract(BaseModel):
    model_config = ConfigDict(
        frozen=True,
        extra="forbid",
        strict=True,
        revalidate_instances="always",
        validate_default=True,
        hide_input_in_errors=True,
    )

    @model_validator(mode="wrap")
    @classmethod
    def exact(cls, value: object, handler: ModelWrapValidatorHandler[Self]) -> Self:
        try:
            if isinstance(value, BaseModel):
                if type(value) is not cls:
                    _fail(UploadParserErrorCode.INVALID_INPUT)
                value = {**vars(value), **(value.__pydantic_extra__ or {})}
            if type(value) is not dict:
                _fail(UploadParserErrorCode.INVALID_INPUT)
            return handler(value)
        except UploadParserError:
            raise
        except Exception:
            raise UploadParserError(UploadParserErrorCode.INVALID_INPUT) from None


class UploadDocument(_Contract):
    document_id: str
    filename: str
    declared_media_type: str | None = None
    content_bytes: Annotated[bytes, Field(strict=True, min_length=1)] = Field(
        repr=False, exclude=True
    )
    source_timestamp: datetime | None = None

    @field_validator("content_bytes", mode="before")
    @classmethod
    def bytes_value(cls, value: object) -> bytes:
        if type(value) is not bytes or not value:
            _fail(UploadParserErrorCode.INVALID_INPUT)
        return value

    @field_validator("document_id", mode="before")
    @classmethod
    def identity(cls, value: object) -> str:
        return _text(value)

    @field_validator("filename", mode="before")
    @classmethod
    def display_name(cls, value: object) -> str:
        text = _text(value)
        if (
            not 1 <= len(text) <= 255
            or text != text.strip()
            or any(ord(c) < 32 or 127 <= ord(c) <= 159 for c in text)
            or any(c in text for c in "/\\:")
            or text in (".", "..")
            or "." not in text
            or not text.rsplit(".", 1)[0].strip(".")
        ):
            _fail(UploadParserErrorCode.INVALID_INPUT)
        if _extension(text) not in _EXTENSIONS:
            _fail(UploadParserErrorCode.UNSUPPORTED_FORMAT)
        return text

    @field_validator("declared_media_type", mode="before")
    @classmethod
    def mime(cls, value: object) -> str | None:
        if value is None:
            return None
        if (
            type(value) is not str
            or len(value) > 255
            or re.fullmatch(r"[A-Za-z0-9!#$&^_.+-]+/[A-Za-z0-9!#$&^_.+-]+", value)
            is None
        ):
            _fail(UploadParserErrorCode.INVALID_INPUT)
        return value

    @field_validator("source_timestamp", mode="before")
    @classmethod
    def timestamp(cls, value: object) -> datetime | None:
        if value is None:
            return None
        if type(value) is not datetime or value.utcoffset() is None:
            _fail(UploadParserErrorCode.INVALID_INPUT)
        return value


class UploadParserLimits(_Contract):
    max_source_bytes: _Positive
    max_text_chars: _Positive
    max_segments: _Positive
    max_pdf_pages: _Positive
    max_pdf_content_stream_bytes: _Positive
    max_docx_entries: _Positive
    max_docx_member_uncompressed_bytes: _Positive
    max_docx_total_uncompressed_bytes: _Positive
    max_docx_compression_ratio: _Positive
    max_csv_rows: _Positive
    max_csv_columns: _Positive
    max_csv_cell_chars: _Positive

    @model_validator(mode="after")
    def relation(self) -> Self:
        if (
            self.max_docx_member_uncompressed_bytes
            > self.max_docx_total_uncompressed_bytes
        ):
            _fail(UploadParserErrorCode.INVALID_INPUT)
        return self


def _check_cancel() -> None:
    task = asyncio.current_task()
    if task is not None and task.cancelling():
        raise asyncio.CancelledError


async def _checkpoint() -> None:
    _check_cancel()
    await asyncio.sleep(0)
    _check_cancel()


def _xml_preflight(content: bytes) -> None:
    def forbidden(*args: object) -> None:
        _fail(UploadParserErrorCode.MALFORMED_DOCUMENT)

    parser = expat.ParserCreate()
    parser.StartDoctypeDeclHandler = forbidden
    parser.EntityDeclHandler = forbidden
    parser.ExternalEntityRefHandler = lambda *args: 0
    parser.Parse(content, True)


def _docx_preflight(data: bytes, limits: UploadParserLimits, *, scan: bool) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            # Admission uses archive spelling before ZipInfo's host normalization.
            for member in members:
                original = member.orig_filename
                if (
                    original.startswith("/")
                    or "\\" in original
                    or ":" in original
                    or ".." in original.split("/")
                    or any(ord(c) < 32 or 127 <= ord(c) <= 159 for c in original)
                    or original != member.filename
                ):
                    _fail(UploadParserErrorCode.MALFORMED_DOCUMENT)
            names = [member.filename for member in members]
            if len(names) != len(set(names)) or not {
                "[Content_Types].xml",
                "word/document.xml",
            } <= set(names):
                _fail(UploadParserErrorCode.MALFORMED_DOCUMENT)
            if len(members) > limits.max_docx_entries:
                _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
            total = 0
            for member in members:
                name = member.filename
                if member.flag_bits & 1:
                    _fail(UploadParserErrorCode.ENCRYPTED_DOCUMENT)
                total += member.file_size
                if (
                    member.file_size > limits.max_docx_member_uncompressed_bytes
                    or total > limits.max_docx_total_uncompressed_bytes
                    or member.file_size
                    > member.compress_size * limits.max_docx_compression_ratio
                ):
                    _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
                if scan and name.lower().endswith((".xml", ".rels")):
                    _check_cancel()
                    with archive.open(member) as stream:
                        content = stream.read(
                            limits.max_docx_member_uncompressed_bytes + 1
                        )
                    _check_cancel()
                    if (
                        len(content) != member.file_size
                        or len(content) > limits.max_docx_member_uncompressed_bytes
                    ):
                        _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
                    _xml_preflight(content)
                    _check_cancel()
    except UploadParserError:
        raise
    except Exception:
        raise UploadParserError(UploadParserErrorCode.MALFORMED_DOCUMENT) from None


def _admit(
    upload: UploadDocument, limits: UploadParserLimits
) -> tuple[UploadFormat, str]:
    if len(upload.content_bytes) > limits.max_source_bytes:
        _fail(UploadParserErrorCode.SOURCE_TOO_LARGE)
    ext = _extension(upload.filename)
    fmt = _EXTENSIONS[ext]
    canonical = _IMAGE_MEDIA[ext] if fmt is UploadFormat.IMAGE else _MEDIA[fmt]
    declared = upload.declared_media_type
    if declared is not None:
        value = declared.lower()
        compatible = value in (canonical, "application/octet-stream")
        compatible |= fmt is UploadFormat.MARKDOWN and value == "text/plain"
        compatible |= fmt is UploadFormat.CODE and (
            value.startswith("text/")
            or value in ("application/json", "application/xml")
        )
        if not compatible:
            _fail(UploadParserErrorCode.MEDIA_TYPE_MISMATCH)
    data = upload.content_bytes
    if fmt is UploadFormat.PDF and not data.startswith(b"%PDF-"):
        _fail(UploadParserErrorCode.MALFORMED_DOCUMENT)
    if fmt is UploadFormat.DOCX:
        _docx_preflight(data, limits, scan=False)
    if fmt is UploadFormat.IMAGE:
        valid = (
            data.startswith(b"\x89PNG\r\n\x1a\n")
            if ext == ".png"
            else data.startswith(b"\xff\xd8\xff")
            if ext in (".jpg", ".jpeg")
            else len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP"
        )
        if not valid:
            _fail(UploadParserErrorCode.MALFORMED_DOCUMENT)
    return fmt, canonical


@dataclass(frozen=True, slots=True)
class ResolvedUploadRoute:
    source: DocumentSource = field(repr=False)
    format: UploadFormat
    parser: DocumentParser = field(repr=False)

    def __post_init__(self) -> None:
        if (
            type(self.source) is not DocumentSource
            or type(self.format) is not UploadFormat
            or not callable(getattr(self.parser, "parse", None))
        ):
            _fail(UploadParserErrorCode.INVALID_INPUT)
        try:
            accepted = DocumentSource.model_validate(self.source)
            object.__setattr__(self, "source", accepted)
        except Exception:
            raise UploadParserError(UploadParserErrorCode.INVALID_INPUT) from None


@contextmanager
def _pdf_boundary(limits: UploadParserLimits) -> Iterator[None]:
    # Synchronous scope only: redact at record creation before any handler sees
    # provider-controlled diagnostics; restore the caller's factory in finally.
    factory = logging.getLogRecordFactory()
    owner = threading.get_ident()

    def private_record(*args: object, **kwargs: object) -> logging.LogRecord:
        record = factory(*args, **kwargs)
        if (
            record.name == "pypdf" or record.name.startswith("pypdf.")
        ) and threading.get_ident() == owner:
            record.msg = "PDF parser diagnostic withheld"
            record.args = ()
            record.exc_info = record.exc_text = record.stack_info = None
        return record

    logging.setLogRecordFactory(private_record)
    bound = limits.max_pdf_content_stream_bytes
    try:
        with pypdf.apply_configuration(
            jbig2dec_binary=None,
            disable_legacy_handling=True,
            maximum_declared_stream_length=bound,
            array_based_stream_maximum_output_length=bound,
            zlib_maximum_output_length=bound,
            lzw_maximum_output_length=bound,
            run_length_maximum_output_length=bound,
            jbig2_maximum_output_length=bound,
            zlib_maximum_recovery_input_length=limits.max_source_bytes,
        ):
            yield
    finally:
        logging.setLogRecordFactory(factory)


class _BuiltinParser:
    def __init__(self, fmt: UploadFormat, limits: UploadParserLimits) -> None:
        self._format = fmt
        self._limits = UploadParserLimits.model_validate(limits)
        self._loop: asyncio.AbstractEventLoop | None = None

    def _pdf_call(self, call: Callable[[], _T]) -> _T:
        _check_cancel()
        with _pdf_boundary(self._limits):
            result = call()
        _check_cancel()
        return result

    def _decode(self, source: DocumentSource, *, meaningful: bool) -> str:
        try:
            text = source.content_bytes.decode("utf-8-sig", errors="strict")
        except UnicodeError:
            raise UploadParserError(UploadParserErrorCode.DECODE_FAILURE) from None
        if "\0" in text:
            _fail(UploadParserErrorCode.DECODE_FAILURE)
        if len(text) > self._limits.max_text_chars:
            _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
        if meaningful and not text.strip():
            _fail(UploadParserErrorCode.NO_EXTRACTABLE_TEXT)
        return text

    async def parse(self, source: DocumentSource) -> ParsedDocument:
        loop = asyncio.get_running_loop()
        if self._loop is not None and self._loop is not loop:
            _fail(UploadParserErrorCode.WRONG_EVENT_LOOP)
        self._loop = loop
        await _checkpoint()
        try:
            if type(source) is not DocumentSource:
                _fail(UploadParserErrorCode.INVALID_INPUT)
            source = DocumentSource.model_validate(source)
            upload = UploadDocument(
                document_id=source.document_id,
                filename=_text(source.metadata["filename"]),
                content_bytes=source.content_bytes,
                source_timestamp=source.source_timestamp,
                declared_media_type=source.media_type,
            )
            fmt, media = _admit(upload, self._limits)
            if (
                fmt is not self._format
                or source.source_type != "upload"
                or source.source_ref != "upload:" + source.document_id
                or source.media_type != media
                or dict(source.metadata)
                != {"filename": upload.filename, "upload_format": fmt.value}
            ):
                _fail(UploadParserErrorCode.INVALID_INPUT)
            segments: list[ParsedSegment] = []
            retained = 0

            def append(
                text: str,
                *,
                page: int | None = None,
                section: str | None = None,
                row: int | None = None,
            ) -> None:
                nonlocal retained
                if not text.strip():
                    return
                if (
                    len(segments) >= self._limits.max_segments
                    or retained + len(text) > self._limits.max_text_chars
                ):
                    _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
                segments.append(
                    ParsedSegment(
                        text=text,
                        path=upload.filename,
                        page=page,
                        section=section,
                        metadata={} if row is None else {"row_number": row},
                    )
                )
                retained += len(text)

            version = "1"
            parser_id = "phoenix." + fmt.value + ".utf8"
            if fmt in (UploadFormat.TEXT, UploadFormat.MARKDOWN, UploadFormat.CODE):
                append(self._decode(source, meaningful=True), section=fmt.value)
            elif fmt is UploadFormat.CSV:
                parser_id = "phoenix.csv.stdlib"
                text = self._decode(source, meaningful=False)
                reader = csv.reader(StringIO(text, newline=""), strict=True)
                row_number = 0
                while True:
                    await _checkpoint()
                    old_limit = csv.field_size_limit()
                    try:
                        csv.field_size_limit(
                            min(self._limits.max_csv_cell_chars, sys.maxsize)
                        )
                        row = next(reader, None)
                    except csv.Error as error:
                        code = (
                            UploadParserErrorCode.LIMIT_EXCEEDED
                            if "field larger than field limit" in str(error)
                            else UploadParserErrorCode.MALFORMED_DOCUMENT
                        )
                        raise UploadParserError(code) from None
                    finally:
                        csv.field_size_limit(old_limit)
                    _check_cancel()
                    if row is None:
                        break
                    row_number += 1
                    if (
                        row_number > self._limits.max_csv_rows
                        or len(row) > self._limits.max_csv_columns
                        or any(
                            len(cell) > self._limits.max_csv_cell_chars for cell in row
                        )
                    ):
                        _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
                    append(
                        json.dumps(row, ensure_ascii=False, separators=(",", ":")),
                        section=f"row:{row_number}",
                        row=row_number,
                    )
            elif fmt is UploadFormat.PDF:
                parser_id, version = "phoenix.pdf.pypdf", pypdf.__version__
                with BytesIO(source.content_bytes) as buffer:
                    reader_pdf = self._pdf_call(
                        lambda: pypdf.PdfReader(buffer, strict=True)
                    )
                    if reader_pdf.is_encrypted:
                        _fail(UploadParserErrorCode.ENCRYPTED_DOCUMENT)
                    count = self._pdf_call(lambda: len(reader_pdf.pages))
                    if count > self._limits.max_pdf_pages:
                        _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
                    for index in range(count):
                        await _checkpoint()
                        page_pdf = self._pdf_call(lambda: reader_pdf.pages[index])
                        contents = self._pdf_call(page_pdf.get_contents)
                        if (
                            contents is not None
                            and len(self._pdf_call(contents.get_data))
                            > self._limits.max_pdf_content_stream_bytes
                        ):
                            _fail(UploadParserErrorCode.LIMIT_EXCEEDED)
                        append(self._pdf_call(page_pdf.extract_text), page=index + 1)
            elif fmt is UploadFormat.DOCX:
                parser_id, version = "phoenix.docx.python-docx", docx.__version__
                await _checkpoint()
                _docx_preflight(source.content_bytes, self._limits, scan=True)
                _check_cancel()
                with BytesIO(source.content_bytes) as buffer:
                    document = docx.Document(buffer)
                    _check_cancel()
                    paragraphs = tables = 0
                    for block in document.iter_inner_content():
                        await _checkpoint()
                        if isinstance(block, Paragraph):
                            paragraphs += 1
                            text = block.text
                            _check_cancel()
                            append(text, section=f"paragraph:{paragraphs}")
                        elif isinstance(block, Table):
                            tables += 1
                            for ordinal, row in enumerate(block.rows, 1):
                                await _checkpoint()
                                text = "\t".join(cell.text for cell in row.cells)
                                _check_cancel()
                                append(text, section=f"table:{tables}/row:{ordinal}")
            else:
                _fail(UploadParserErrorCode.PARSER_UNAVAILABLE)
            await _checkpoint()
            if not segments:
                _fail(UploadParserErrorCode.NO_EXTRACTABLE_TEXT)
            result = ParsedDocument(
                document_id=source.document_id,
                parser_id=parser_id,
                parser_version=version,
                segments=tuple(segments),
            )
            _check_cancel()
            return result
        except pypdf.errors.LimitReachedError:
            _check_cancel()
            raise UploadParserError(UploadParserErrorCode.LIMIT_EXCEEDED) from None
        except UploadParserError:
            _check_cancel()
            raise
        except Exception:
            _check_cancel()
            raise UploadParserError(UploadParserErrorCode.MALFORMED_DOCUMENT) from None


class UploadRouter:
    """Deterministic admission/routing, without parsing or ownership of adapters."""

    def __init__(
        self, limits: UploadParserLimits, *, image_parser: DocumentParser | None = None
    ) -> None:
        if type(limits) is not UploadParserLimits:
            _fail(UploadParserErrorCode.INVALID_INPUT)
        self._limits = UploadParserLimits.model_validate(limits)
        if image_parser is not None and not callable(
            getattr(image_parser, "parse", None)
        ):
            _fail(UploadParserErrorCode.INVALID_INPUT)
        self._image_parser = image_parser
        self._parsers = {
            fmt: _BuiltinParser(fmt, self._limits)
            for fmt in UploadFormat
            if fmt is not UploadFormat.IMAGE
        }

    def resolve(self, upload: UploadDocument) -> ResolvedUploadRoute:
        if type(upload) is not UploadDocument:
            _fail(UploadParserErrorCode.INVALID_INPUT)
        upload = UploadDocument.model_validate(upload)
        fmt, media = _admit(upload, self._limits)
        parser = self._image_parser if fmt is UploadFormat.IMAGE else self._parsers[fmt]
        if parser is None:
            _fail(UploadParserErrorCode.PARSER_UNAVAILABLE)
        try:
            source = DocumentSource(
                document_id=upload.document_id,
                source_type="upload",
                source_ref="upload:" + upload.document_id,
                media_type=media,
                content_bytes=upload.content_bytes,
                source_timestamp=upload.source_timestamp,
                metadata={"filename": upload.filename, "upload_format": fmt.value},
            )
            if (
                source.document_id != upload.document_id
                or source.source_ref != "upload:" + upload.document_id
            ):
                _fail(UploadParserErrorCode.INVALID_INPUT)
            return ResolvedUploadRoute(source=source, format=fmt, parser=parser)
        except UploadParserError:
            raise
        except Exception:
            raise UploadParserError(UploadParserErrorCode.INVALID_INPUT) from None
